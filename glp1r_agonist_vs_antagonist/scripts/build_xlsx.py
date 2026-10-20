# -*- coding: utf-8 -*-
import json, sys, csv, os
sys.path.insert(0,'.')
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import data_sheets as D

OUT="/home/user/ForCluadeCloudSessions/glp1r_agonist_vs_antagonist"
os.makedirs(OUT,exist_ok=True)
wb=Workbook()
HFILL=PatternFill("solid",fgColor="1F3A5F"); HFONT=Font(bold=True,color="FFFFFF",name="Calibri",size=11)
thin=Side(style="thin",color="D0D0D0"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
def sheet(title, header, rows, widths, freeze="A2", csvname=None):
    ws=wb.create_sheet(title)
    ws.append(header)
    for r in rows: ws.append(list(r))
    for c in range(1,len(header)+1):
        cell=ws.cell(row=1,column=c); cell.fill=HFILL; cell.font=HFONT; cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[get_column_letter(c)].width=widths[c-1]
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row,max_col=len(header)):
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical="top"); cell.border=BORDER
    ws.freeze_panes=freeze; ws.auto_filter.ref=ws.dimensions
    if csvname:
        with open(f"{OUT}/{csvname}","w",newline="",encoding="utf-8-sig") as f:
            w=csv.writer(f); w.writerow(header); w.writerows(rows)
    return ws

# ---------- README
ws=wb.active; ws.title="README"
readme=[
 ["Что это","Литобзор: чем отличается связывание агонистов и антагонистов GLP-1R и как это учитывать при дизайне (Challenge 3 конкурса Anthropic × Adaptyv: «Design a GLP1R agonist»)."],
 ["Дата сборки","2026-10-20"],
 ["Как читать","Листы: 1) Сравнение — матрица агонист/антагонист; 2) Лиганды — ключевые молекулы с числами; 3) SAR пептида — позиции GLP-1 7–36; 3b) Остатки рецептора — общие и агонист-специфичные; 4) PDB реестр — все 80 записей RCSB с UniProt P43220; 5) Чек-лист дизайна; 6–9) Расчёты автора; 10) Источники."],
 ["Происхождение чисел","Цифры из статей — со ссылкой на источник в строке. Числа с пометкой «расчёт» посчитаны автором на структурах RCSB (Biopython ShrakeRupley, зонд 1.4 Å; ECD = остатки 29–137; суперпозиция по ядру TMD). Нумерация остатков рецептора — UniProt P43220 (сигнальный пептид 1–23)."],
 ["Ограничения","Все структуры Gs/Nb35-стабилизированы и описывают активный комплекс, а не энергетику перехода. Часть записей PDB 2025–2026 без публикации (в реестре помечено). Условия теста конкурса (линия клеток, экспрессия, одна доза или кривая) организаторами не раскрыты."],
 ["Конкурс","Страница задачи: https://proteinbase.com/competitions/anthropic-adaptyv-2026/challenges/glp1r ; приём дизайнов до 25 октября 2026, 23:59 AoE."],
]
for r in readme: ws.append(r)
ws.column_dimensions["A"].width=24; ws.column_dimensions["B"].width=140
for row in ws.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    row[0].font=Font(bold=True)

# ---------- 1. Comparison matrix
hdr=["Признак","Пептидный агонист (GLP-1, Ex4, семаглутид, тирзепатид)","Непептидный агонист (PF-06882961, LY3502970, TT-OAD2, CHU-128)","PAM / ago-PAM (BETP, LSN3318839, compound 2)","ECD-антагонист (Ex(9–39), Fab 3F52, TB-222-023, de novo миниbелки)","NAM в TMD (PF-06372222, NNC0640)"]
rows=[
 ["Основной сайт","ECD-борозда (C-половина, остатки 24–36) + TMD-карман (N-половина 7–17) + ECL1/ECL2","Верхний карман TMD; у LY3502970 — ECD, ECL2, TM1/2/3/7; у TT-OAD2 новый карман у ECL3/TM6/TM7","Цитоплазматический конец TM6 (compound 2 ковалентно на C347), липид-обращённая поверхность","ECD-борозда (+ у Fab 3F52 стык ECD–TM1, остатки 137–145)","Внеспиральный карман снаружи TM5–7 у цитоплазматической половины"],
 ["Вклад ECD в связывание (скрыто, Å², расчёт)","≈ 550–690 (GLP-1 6X18: 651)","≈ 150–200 (PF-06882961 199, LY3502970 170, Boc5 146)","—","Основной вклад (Fab 3F52 ≈ 970 Å² на ECD; Ex(9–39)/ECD 688 Å² на лиганде)","Нет"],
 ["Вклад TMD/ECL в связывание (Å², расчёт)","≈ 1100–1300 (GLP-1: 1265; семаглутид 1152; тирзепатид 1108)","≈ 600–800","—","Нет (N-конец отсутствует) / лишь стерика у входа","—"],
 ["N-концевой модуль (H7–F12) в кармане","Да; 7–17 ≈ 890 Å² только на TMD/ECL, нулевой контакт с ECD","Нет глубокой вставки; центр лиганда на 13–15 Å от Cα H7 (пептиды 22–23 Å)","Не требуется","Нет: Ex(9–39) без остатков 1–8; Fab/миниbелки не заходят в карман","Не относится"],
 ["Положение ECD","Открыто (центр ECD на 16 Å от закрытой апо-формы)","Промежуточное: 5–7 Å от пептид-связанного, 10–11 Å от закрытого","Меняется (в 7DUR N-конец ECD T29–E34 лежит у места пептида)","ECD связан лигандом; полный рецептор без агониста — закрытая форма (6LN2)","Закрытая (6LN2)"],
 ["TM6 (Cα 346)","Наружу на 18–19 Å от неактивной формы","То же (18.0–18.8 Å)","То же","Не переводит в активное состояние","Удерживает внутриклеточный конец TM6"],
 ["Эффект","Полный агонист; смещение (bias) зависит от N-концевой динамики","Часто частичные и G-смещённые","Усиливают пептид (PAM) или сами активируют (ago-PAM)","Конкурентный/нейтральный антагонист","Негативная аллостерическая модуляция"],
 ["Связь аффинность ↔ эффективность","У GLP-1 сильная; у экзендина разнесены по доменам","Аффинность низкая, эффективность есть","—","Аффинность высокая (3–6 нМ), эффективность отсутствует","—"],
 ["Чувствительность к тест-системе","Низкая","Частичные агонисты зависят от экспрессии","—","Метаболит GLP-1(9–36): агонист при высокой экспрессии, антагонист при низкой (Knudsen 2012)","—"],
 ["Примеры PDB","6X18, 7KI0, 7FIM, 7LLL, 6B3J, 9BYO","6X1A, 6X19, 6XOX, 6ORV, 7S15, 7X8R","6VCB, 7DUQ, 7DUR, 9IVM, 9NTU","3C59, 3C5T, 5E94","5VEW, 5VEX, 6KJV, 6LN2"],
 ["Ключевые источники","Zhang 2017 (10.1038/nature22394); Zhang 2020 (10.1016/j.molcel.2020.09.020); Montrose-Rafizadeh 1997; Adelhorst 1994","Kawai 2020 (10.1073/pnas.2014879117); Zhao 2020 (10.1038/s41586-019-1902-z); Zhang 2020","Bueno 2020 (10.1038/s41589-020-0589-7); Cong 2021 (10.1038/s41467-021-24058-z)","Runge 2008 (10.1074/jbc.M708740200); Hennen 2016 (10.1038/srep26236); Peterson 2023; Baker 2025 (10.1101/2025.03.23.644666)","Song 2017 (10.1038/nature22378); Wu 2020 (10.1038/s41467-020-14934-5)"],
]
sheet("1 Сравнение",hdr,rows,[26,44,44,38,44,38],csvname="1_comparison_matrix.csv")

# ---------- 2. Ligands
hdr=["Лиганд","Тип","Последовательность / химия","Сайт на ECD","Контакт с TMD","Эффект","Ключевые числа","PDB","Источник"]
L=[
 ["GLP-1(7–36)NH2","пептид, агонист","HAEGTFTSDVSSYLEGQAAKEFIAWLVKGR","Борозда, остатки 24–36 (≈ 650 Å² на ECD в 6X18)","H7–F12 в кармане; ECL1, ECL2","Полный агонист","На изолированном ECD IC50 > 500 нМ; Ala-замены: F28A ×1300, I29A ×93, L32A ×17","6X18, 3IOL","Underwood 2010; Zhang 2020; Adelhorst 1994"],
 ["GLP-1(9–36)NH2","метаболит DPP-4","EGTFTSDVSSYLEGQAAKEFIAWLVKGR","Борозда","Без H7–A8","Слабый частичный агонист; при низкой экспрессии — слабый антагонист","—","9IVG","Knudsen 2012; Acta Pharm Sin B 2024"],
 ["Exendin-4","пептид, агонист","HGEGTFTSDLSKQMEEEAVRLFIEWLKNGGPSSGAPPPS","Борозда; аффинность выше за счёт спиральности остатков Leu10–Gly30; Trp-cage в связывании не участвует","Да","Полный агонист","IC50 на ECD ≈ 6 нМ","7LLL, 3C59","Runge 2007; Runge 2008"],
 ["Exendin-(9–39)NH2 (avexitide)","пептид, конкурентный антагонист","DLSKQMEEEAVRLFIEWLKNGGPSSGAPPPS","Борозда (те же ≈ 14 из 16–18 остатков, что у GLP-1)","Нет (нет остатков 1–8)","Антагонист; клинически — гипогликемия после бариатрической операции","IC50 на ECD ≈ 6 нМ; t½ ≈ 30 мин в/в; 44 клинических исследования у людей; ChEMBL: фаза 3","3C5T, 3C59","Runge 2008; Gasbjerg 2021"],
 ["Exendin-(3–39)","пептид, антагонист","—","Борозда","Нет","Антагонист в 4–10× мощнее Ex(9–39)","Усечение на 3–7 а.о. даёт сильнейшие антагонисты","—","Montrose-Rafizadeh 1997"],
 ["Семаглутид","пептид, агонист (Aib8, C18-диацид)","—","Борозда","Да","Полный агонист","—","7KI0","Zhang X 2021 Cell Rep"],
 ["Тирзепатид","пептид, дуальный агонист","—","Борозда","Да","Смещён к cAMP на GLP-1R","—","7FIM, 7RGP","Willard 2020; Zhao F 2022; Sun B 2022"],
 ["Exendin-P5","пептид, G-смещённый","—","Борозда","Да","G-смещённый","Угол α5 Gαs на 6° отличается","6B3J","Liang 2018"],
 ["Fab 3F52","антитело (Fab), антагонист","—","17 остатков ECD, ≈ 970 Å²; перекрывает 7 из 16 остатков борозды и стык ECD–TM1 (137–145) (расчёт)","Нет","Нейтральный конкурентный антагонист; блокирует и короткий пептидный агонист, но не аллостерический малый агонист","Связывание IC50 ≈ 3 нМ; cAMP IC50 ≈ 10 нМ; агонизма нет до 1 мкМ; не связывает мышь, крысу, свинью","5E94","Hennen 2016"],
 ["TB-222-023","IgG, антагонист","—","ECD (фаг-дисплей на клетках, GPCR-ориентированная библиотека)","Нет","Антагонист","Эффективнее avexitide; работает в островках пациента с гиперинсулинизмом","—","Peterson 2023; Liu 2021"],
 ["dGl1_024 / mGl1_008","de novo миниbелки, антагонисты","RFdiffusion / MetaGen + ProteinMPNN","ECD, хотспот «Leu95» = Leu123 по UniProt (порядковый номер в усечённой цепи 5VAI/6X18/7DUQ; наша реконструкция, авторами не подтверждена)","Нет","Антагонисты","IC50 61 / 39 нМ; Kd 27 / 5.3 нМ; 49 экспрессированных из ≈ 17 000 дизайнов","—","Baker 2025 (препринт; Nature 2026)"],
 ["Антитело с GLP-1 на лёгкой цепи","агонист","GLP-1, слитый с лёгкой цепью нефункционального GLP-1R-антитела","ECD (антитело)","Да, за счёт слитого пептида","Агонист","Из 13 антител большинство антагонисты","—","Liu 2021"],
 ["PF-06882961 (данаглипрон)","малая молекула, агонист","—","Контакт с ECD ≈ 199 Å²","Верхний карман; сайт пересекается с GLP-1","Агонист, близко имитирует GLP-1","Центр лиганда на 15.0 Å от Cα H7","6X1A, 7LCI","Zhang 2020 Mol Cell; Saxena 2021"],
 ["LY3502970 (орфорглипрон)","малая молекула, частичный агонист","—","ECD, ECL2; W33 ECD (приматоспецифичный)","TM1/2/3/7","Частичный, смещён к G-белку","ECD 170 Å² / TMD 741 Å²","6XOX","Kawai 2020"],
 ["TT-OAD2","малая молекула, G-смещённый","—","—","Новый карман; ECL3/TM6/TM7 перестроены без прямого контакта в глубине","G-смещённый; иная кинетика","—","6ORV","Zhao 2020"],
 ["CHU-128","малая молекула","—","—","Более открытая внеклеточная сторона","Агонист с особым режимом","—","6X19","Zhang 2020 Mol Cell"],
 ["Compound 2","ago-PAM, ковалентный","—","—","C347 на цитоплазматическом конце TM6","Усиливает агонист и сам активирует","В 7DUR N-конец ECD (T29–E34) лежит в 3.8–5.3 Å от места остатков 7–15 пептида (расчёт)","7DUQ, 7DUR, 7EVM","Cong 2021"],
 ["BETP / LSN3318839","PAM","—","—","Цитоплазматическая область TM6","Зависимый от зонда позитивный аллостеризм","LSN3318839 + GLP-1(9–36): усиление сигнала метаболита","6VCB, 9IVM","Bueno 2020; Acta Pharm Sin B 2024"],
 ["PF-06372222 / NNC0640","NAM (малые молекулы)","—","—","Карман снаружи TM5–7","Негативный аллостерический модулятор","2.7 / 3.0 Å","5VEW, 5VEX","Song 2017"],
 ["EvoBind-пептиды (18-мер, 31-мер, 15-мер)","циклические пептиды, de novo","18-мер: LCWYGYNDPQAPAMVNWY (по методам препринта; соответствие длинам выведено)","Предсказана «поверхность активации», ECD/TMD не разделены","Предсказано глубоко в кармане","Агонисты cAMP","18-мер EC50 32 нМ (лираглутид 2.4 нМ); 31-мер 5.0 мкМ; 15-мер 11.8 мкМ","— (модель)","Li 2025, препринт"],
]
sheet("2 Лиганды",hdr,L,[30,26,44,40,36,36,44,14,36],csvname="2_ligands.csv")

# ---------- 3. Peptide SAR
fig1=json.load(open("work/fig1_data.json"))
be={r:b for r,b in zip(fig1["res"],fig1["bur_ecd"])}; bt={r:b for r,b in zip(fig1["res"],fig1["bur_tmd"])}
contacts={7:"Q234 V237 W306 I309 R310 I313",8:"L384 E387 L388",9:"Y152 R190 L388",10:"N300 W306",11:"W306 D372 R380 L384",12:"L141 L144 Y148 L388",13:"K197 F230 M233 T298",14:"T298 R299 N300",15:"R380 L384",16:"L201",17:"Y205 T298 R299",18:"R299",19:"E138",20:"Y205"}
rows=[]
for pos in range(7,37):
    i=pos-7
    dom = "TMD/ECL" if pos<=17 else ("смесь" if pos<=23 else "ECD")
    rows.append([pos, D.GLP1[i], D.EX4[i], (D.GCG[i] if i<len(D.GCG) else ""), D.GIP[i], dom, be.get(pos,0), bt.get(pos,0), contacts.get(pos,""), D.POS_NOTES.get(pos,"")])
sheet("3 SAR пептида",["Позиция GLP-1","GLP-1","Ex4 (экв.)","Глюкагон (экв.)","GIP (экв.)","Домен контакта в 6X18","Скрыто ECD, Å² (расчёт)","Скрыто TMD/ECL, Å² (расчёт)","Контакты ≤ 4 Å с TMD/ECL","Роль и источники"],rows,[10,8,10,12,10,14,14,14,30,80],csvname="3_peptide_SAR_by_position.csv")

# ---------- 3b. Hotspots
sheet("3b Остатки рецептора",["Остаток (UniProt P43220)","Элемент","Контакт с GLP-1 в 6X18 (≤4 Å)","Данные мутагенеза / функции","Класс","Источник"],D.HOTSPOTS,[34,22,36,80,24,30],csvname="3b_receptor_hotspots.csv")

# ---------- 4. PDB registry
det=json.load(open("work/glp1r_pdb_details.json"))["data"]["entries"]
rows=[]
for e in det:
    pid=e["rcsb_id"]; cit=e.get("rcsb_primary_citation") or {}
    res=((e.get("rcsb_entry_info") or {}).get("resolution_combined") or [None])[0]
    date=((e.get("rcsb_accession_info") or {}).get("initial_release_date") or "")[:10]
    cls=D.PDB_CLASS[pid]
    rows.append([pid,date,(e["exptl"][0]["method"] if e.get("exptl") else ""),res,cls[0],cls[1],cls[2],e["struct"]["title"],cit.get("pdbx_database_id_DOI") or "",cit.get("pdbx_database_id_PubMed") or "",f'{cit.get("journal_abbrev") or ""} {cit.get("year") or ""}'.strip()])
rows.sort(key=lambda r:r[1])
sheet("4 PDB реестр",["PDB","Выпуск","Метод","Разрешение, Å","Класс (наша классификация)","Лиганд","Заметка","Название записи RCSB","DOI","PMID","Журнал, год"],rows,[8,11,18,10,36,34,48,60,28,10,18],csvname="4_pdb_registry_80_entries.csv")

# ---------- 5. Design checklist
sheet("5 Чек-лист дизайна",["#","Принцип","Опора (источник / расчёт)","Что делать в дизайне","Уверенность","Как проверить"],D.DESIGN_ROWS,[4,34,80,60,20,34],csvname="5_design_checklist.csv")

# ---------- 6. Calc: GLP-1 burial
rows=[[r,l,be[r],bt[r]] for r,l in zip(fig1["res"],fig1["label"])]
sheet("6 Расчёт - GLP-1 в 6X18",["Позиция","Остаток","Скрыто ECD, Å²","Скрыто TMD/ECL, Å²"],rows,[10,10,16,18],csvname="6_calc_GLP1_burial_6X18.csv")

# ---------- 7. Calc: state comparison
st=json.load(open("work/states.json"))
names={"6X18":"GLP-1(7–36)","7KI0":"Семаглутид","7FIM":"Тирзепатид","7LLL":"Exendin-4","6B3J":"Exendin-P5 (G-смещённый)","9IVG":"GLP-1(9–36)","6VCB":"GLP-1 + PAM LSN3318839 (лиганд = пептид + PAM)","6X1A":"PF-06882961","6X19":"CHU-128","6XOX":"LY3502970","7S15":"Малая молекула Pfizer","7RG9":"Апо, Gs","8WG7":"Пептид-свободный, Gs","7C2E":"Малая молекула FFR","7X8R":"Boc5"}
rows=[]
for x in st:
    rows.append([x["pid"],names.get(x["pid"],""),x.get("core"),x.get("rmsd"),x.get("TM6_346_to_active"),x.get("TM6_346_to_inactive"),x.get("ECL1_215_to_active"),x.get("ECL1_215_to_inactive"),x.get("ECD_to_active"),x.get("ECD_to_inactive"),x.get("lig_free_SASA"),x.get("bur_ECD"),x.get("bur_TMD"),x.get("cent_to_H7")])
sheet("7 Расчёт - состояния",["PDB","Лиганд","Остатков в ядре TMD","RMSD ядра, Å","Cα346 до 6X18, Å","Cα346 до 6LN2, Å","Cα215 до 6X18, Å","Cα215 до 6LN2, Å","Центр ECD до 6X18, Å","Центр ECD до 6LN2, Å","SASA лиганда свободного, Å²","Скрыто ECD, Å²","Скрыто TMD/ECL, Å²","Центр лиганда от Cα H7, Å"],rows,[8,34,10,10,12,12,12,12,12,12,14,12,14,14],csvname="7_calc_states.csv")

# ---------- 8. Calc: ECD groove
groove=[["Ex4(9–39), 3C5T (ECD)","28 30 31 32 35 36 39 67 68 69 88 90 91 118 121 123 127 128","18","688 (на лиганде)"],
        ["GLP-1, 3IOL (изолированный ECD)","29 30 32 35 36 39 67 68 69 88 89 90 91 121 123 128","16","571 (на лиганде)"],
        ["GLP-1, 6X18 (полноразмерный)","29 30 31 32 35 36 39 43 67 68 69 88 90 91 121 123 128 137","18","651 (на лиганде)"],
        ["Fab 3F52, 5E94 (эпитоп)","29 30 31 32 33 35 36 39 40 43 68 137 138 141 142 144 145","17","≈ 970 (на ECD)"],
        ["Совпадение: Ex4(9–39) ∩ GLP-1(3IOL)","14 общих остатков","",""],
        ["Совпадение: GLP-1 (3IOL) ∩ GLP-1 (6X18)","15 общих остатков","",""],
        ["Совпадение: Fab 3F52 ∩ GLP-1 (3IOL)","29 30 32 35 36 39 68 — 7 из 16 (44%)","",""],
        ["Доступность борозды в закрытой форме (6LN2)","994 Å² в контексте / 1125 Å² на изолированном ECD = 0.88","",""],
        ["Доступность борозды в 6X18 (рецептор без пептида)","1022 Å² / 1152 Å² = 0.89","",""]]
sheet("8 Расчёт - борозда ECD",["Структура / сравнение","Остатки ECD (UniProt), скрытые > 10 Å²","Число остатков","Скрыто, Å²"],groove,[46,80,14,22],csvname="8_calc_ECD_groove.csv")

# ---------- 9. Calc: inactive -> active
rows=[["Центр ECD, сдвиг 6LN2 → 6X18","16.1 Å"],["Средний сдвиг Cα ECD (29–137)","22.6 Å (макс. 34.1)"],["Cα остатка 346 (внутриклеточный конец TM6)","18.7 Å"],["ECL1 / верх TM2–TM3 (198–225)","до 15.5 Å (остаток 215)"],["TM6–TM7 (344–390)","до 18.7 Å"],["Ядро для суперпозиции","145 Cα TMD, RMSD 1.24 Å (итеративное отбрасывание выбросов > 3 Å)"]]
sheet("9 Расчёт - неакт. в акт.",["Показатель","Значение"],rows,[56,70],csvname="9_calc_inactive_to_active.csv")

# ---------- 10. Sources
S=[
 ["Runge S et al. 2008, J Biol Chem — структура ECD с Ex4(9–39)","10.1074/jbc.M708740200","18287102"],
 ["Underwood CR et al. 2010, J Biol Chem — структура ECD с GLP-1","10.1074/jbc.M109.033829","19861722"],
 ["López de Maturana R et al. 2003, J Biol Chem — изолированный ECD связывает экзендины сильнее GLP-1","10.1074/jbc.M212147200","12524435"],
 ["Montrose-Rafizadeh C et al. 1997, J Biol Chem — N-усечения экзендина и GLP-1","10.1074/jbc.272.34.21201","9261127"],
 ["Adelhorst K et al. 1994, J Biol Chem — Ala-скан GLP-1","","8119974"],
 ["Gallwitz B et al. 1994, Eur J Biochem — Ala-скан GLP-1","10.1111/j.1432-1033.1994.1151b.x","7957206"],
 ["Hjorth SA et al. 1994, J Biol Chem — глюкагон/GLP-1, химеры","","7527026"],
 ["Al-Sabah S, Donnelly D 2003, Br J Pharmacol — модель связывания","10.1038/sj.bjp.0705453","12970080"],
 ["Al-Sabah S, Donnelly D 2003, FEBS Lett — K288","10.1016/s0014-5793(03)01043-3","14572647"],
 ["Runge S et al. 2007, Biochemistry — спиральность лигандов и аффинность к ECD","10.1021/bi062309m","17444618"],
 ["Zhang Y et al. 2017, Nature — крио-ЭМ GLP-1R–Gs","10.1038/nature22394","28538729"],
 ["Jazayeri A et al. 2017, Nature — кристалл полноразмерного GLP-1R с пептидом","10.1038/nature22800","28562585"],
 ["Song G et al. 2017, Nature — TMD с NAM","10.1038/nature22378","28514449"],
 ["Liang YL et al. 2018, Nature — смещённый агонист Exendin-P5","10.1038/nature25773","29466332"],
 ["Zhang X et al. 2020, Mol Cell — 6X18 и непептидные агонисты","10.1016/j.molcel.2020.09.020","33027691"],
 ["Zhao P et al. 2020, Nature — TT-OAD2","10.1038/s41586-019-1902-z","31915381"],
 ["Wu F et al. 2020, Nat Commun — полноразмерный GLP-1R без лиганда (закрытый ECD)","10.1038/s41467-020-14934-5","32152292"],
 ["Kawai T et al. 2020, PNAS — LY3502970","10.1073/pnas.2014879117","33177239"],
 ["Bueno AB et al. 2020, Nat Chem Biol — PAM LSN3318839","10.1038/s41589-020-0589-7","32690941"],
 ["Cong Z et al. 2021, Nat Commun — ago-PAM compound 2","10.1038/s41467-021-24058-z","34145245"],
 ["Zhang X et al. 2021, Cell Rep — семаглутид и таспоглутид","10.1016/j.celrep.2021.109374","34260945"],
 ["Zhao F et al. 2022, Nat Commun — тирзепатид","10.1038/s41467-022-28683-0","35217653"],
 ["Sun B et al. 2022, PNAS — тирзепатид","10.1073/pnas.2116506119","35333651"],
 ["Deganutti G et al. 2022, Nat Commun — динамика N-конца и кинетика Gs","10.1038/s41467-021-27760-0","35013280"],
 ["Cary BP et al. 2022, Nat Chem Biol — состояния рецептора и склонность N-конца к спирали","10.1038/s41589-021-00945-w","34937906"],
 ["Oddo A et al. 2018, Biochemistry — ограниченные N-концы","10.1021/acs.biochem.8b00105","29877701"],
 ["Hennen S et al. 2016, Sci Rep — Fab 3F52","10.1038/srep26236","27196125"],
 ["Peterson SM et al. 2023, Diabetes — антитело-антагонист TB-222-023","10.2337/db22-1039","37358194"],
 ["Liu Q et al. 2021, mAbs — функциональные антитела к GLP-1R","10.1080/19420862.2021.1893425","33706686"],
 ["Gasbjerg LS et al. 2021, Diabetes Obes Metab — Ex(9–39) в клинике","10.1111/dom.14507","34351033"],
 ["Knudsen LB et al. 2012, Regul Pept — экспрессия рецептора и функция","10.1016/j.regpep.2011.12.006","22252224"],
 ["Yin Y et al. 2016, Cell Discov — собственный агонист ECD","10.1038/celldisc.2016.42","27917297"],
 ["Fortin JP et al. 2009, PNAS — мембранно-привязанные лиганды","10.1073/pnas.0900149106","19416829"],
 ["Cheloha RW et al. 2020, Nat Commun — CLAMP: нанотело + пептид","10.1038/s41467-020-15884-8","32350260"],
 ["Willard FS et al. 2020, JCI Insight — смещение тирзепатида","10.1172/jci.insight.140532","32730231"],
 ["van der Velden WJC et al. 2021, ACS Pharmacol Transl Sci — GLP-1 Val8","10.1021/acsptsci.0c00193","33615180"],
 ["Lei S et al. 2018, J Biol Chem — домены смещённого агонизма","10.1074/jbc.RA118.003278","29717000"],
 ["Jones B et al. 2018, Nat Commun — трафик рецептора","10.1038/s41467-018-03941-2","29686402"],
 ["Acta Pharm Sin B 2024 — GLP-1(9–36) + LSN3318839","10.1016/j.apsb.2024.09.002","39664415"],
 ["Baker lab, bioRxiv 2025 / Nature 2026 — de novo миниbелки для GPCR","10.1101/2025.03.23.644666 ; 10.1038/s41586-026-10656-8",""],
 ["Li Q et al. 2025, bioRxiv — EvoBind, циклические пептиды GCGR/GLP1R","10.1101/2025.06.06.658268",""],
 ["Wei T et al. 2025, Adv Sci — AI-дизайн GLP-1RA (ProteinMPNN, 13 фиксированных сайтов)","10.1002/advs.202507044","40787887"],
 ["Fan S et al. 2025, JACS — мультициклические пептиды, 9J1P","10.1021/jacs.4c12808","39688263"],
 ["Ma Y et al. 2020, Sci Adv — sdAb-агонист рецептора апелина","10.1126/sciadv.aax7379","31998837"],
 ["Fontaine T et al. 2024, Nat Commun — нанотело-агонист MC4R","10.1038/s41467-024-50827-7","39353917"],
 ["Chai-2 preprint, bioRxiv 2025 — функциональные агонисты GPCR (класс A)","10.1101/2025.11.29.691346",""],
 ["Skiba MA et al. 2026, Mol Pharmacol — нанотела и GPCR","10.1016/j.molpha.2026.100110","41795245"],
 ["White AD et al. 2026, Mol Pharmacol — ECD и сигналинг класса B1","10.1016/j.molpha.2026.100111","41795246"],
 ["Zhou Q et al. 2025, Trends Pharmacol Sci — обзор структурной фармакологии GLP-1R","10.1016/j.tips.2025.03.003","40221226"],
 ["Liang YL et al. 2020, Mol Cell — пептидное связывание класса B","10.1016/j.molcel.2020.01.012","32004469"],
 ["Coopman K et al. 2011, Mol Endocrinol — мутагенез TMD GLP-1R","10.1210/me.2011-1160","21868452"],
 ["Xiao Q et al. 2000, J Mol Endocrinol — детерминанты связывания GLP-1R","10.1677/jme.0.0250321","11116211"],
 ["Wootten D et al. 2013, PNAS — полярные TM-остатки и смещённый агонизм","10.1073/pnas.1221585110","23479653"],
 ["Mann RJ et al. 2010, Peptides — Cys226–Cys296","10.1016/j.peptides.2010.09.015","20869417"],
 ["Страница Challenge 3 (Proteinbase)","https://proteinbase.com/competitions/anthropic-adaptyv-2026/challenges/glp1r",""],
 ["Страница выбора дизайнов (Proteinbase)","https://proteinbase.com/competitions/anthropic-adaptyv-2026/selection",""],
 ["Adaptyv: критерии новизны","https://www.adaptyvbio.com/blog/novelty",""],
 ["InvivoGen: HEK293 CRE-SEAP GLP-1 биотест (для ориентира по формату теста)","https://www.invivogen.com/incretin-bioassay-glp-1-reporter-cell-line",""],
 ["RCSB PDB (Search API, Data API), UniProt P43220","https://www.rcsb.org",""],
]
sheet("10 Источники",["Источник","DOI / URL","PMID"],S,[100,60,12],csvname="10_sources.csv")

wb.save(f"{OUT}/GLP1R_agonist_vs_antagonist.xlsx")
print("saved", len(wb.sheetnames), wb.sheetnames)
